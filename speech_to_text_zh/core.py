import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from .constants import APP_DIR, SUPPORTED_EXTENSIONS

def now_stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def clean_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()
    return cleaned or "audio"


def fmt_time(seconds: float) -> str:
    seconds = max(0.0, float(seconds or 0))
    total = int(round(seconds))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def fmt_srt_time(seconds: float) -> str:
    seconds = max(0.0, float(seconds or 0))
    whole = int(seconds)
    ms = int(round((seconds - whole) * 1000))
    if ms >= 1000:
        whole += ms // 1000
        ms = ms % 1000
    h, rem = divmod(whole, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def strip_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


def is_supported_audio(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def parse_audio_paths_text(raw: str) -> List[Path]:
    chunks = [item.strip().strip('"') for item in re.split(r"[;\n]+", raw or "")]
    return [Path(item) for item in chunks if item]


def describe_audio_paths(paths: Sequence[Path]) -> str:
    if not paths:
        return ""
    if len(paths) == 1:
        return str(paths[0])
    names = "；".join(path.name for path in paths[:4])
    suffix = " 等" if len(paths) > 4 else ""
    return f"已选择 {len(paths)} 个文件：{names}{suffix}"


def fmt_duration_seconds(seconds: Optional[float]) -> str:
    if seconds is None:
        return "未知"
    seconds = max(0.0, float(seconds))
    h, rem = divmod(int(round(seconds)), 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}小时{m:02d}分{s:02d}秒"
    if m:
        return f"{m}分{s:02d}秒"
    return f"{s}秒"


def parse_duration_seconds(duration: Optional[str]) -> Optional[float]:
    if not duration:
        return None
    match = re.match(r"(\d+):(\d+):(\d+(?:\.\d+)?)", duration)
    if not match:
        return None
    hours, minutes, seconds = match.groups()
    return int(hours) * 3600 + int(minutes) * 60 + float(seconds)


def segment_seconds(segment: dict) -> float:
    return max(0.0, float(segment.get("end", 0) or 0) - float(segment.get("start", 0) or 0))


def build_quality_summary(
    reliable: Sequence[dict],
    candidates: Sequence[dict],
    discarded: Sequence[dict],
    duration: Optional[str],
) -> Dict[str, Any]:
    reliable_seconds = sum(segment_seconds(seg) for seg in reliable)
    candidate_seconds = sum(segment_seconds(seg) for seg in candidates)
    discarded_seconds = sum(segment_seconds(seg) for seg in discarded)
    audio_seconds = parse_duration_seconds(duration)

    if not reliable:
        label = "需要人工核听"
    elif candidate_seconds + discarded_seconds > reliable_seconds * 0.6:
        label = "可用但需重点复核"
    elif candidates or discarded:
        label = "可用，建议抽查"
    else:
        label = "较干净"

    reliable_ratio = None if not audio_seconds else round(reliable_seconds / audio_seconds, 4)
    return {
        "quality_label": label,
        "audio_seconds": audio_seconds,
        "reliable_count": len(reliable),
        "candidate_count": len(candidates),
        "discarded_count": len(discarded),
        "reliable_seconds": round(reliable_seconds, 2),
        "candidate_seconds": round(candidate_seconds, 2),
        "discarded_seconds": round(discarded_seconds, 2),
        "reliable_ratio_of_audio": reliable_ratio,
    }


def format_quality_summary(summary: Dict[str, Any]) -> str:
    return (
        f"{summary['quality_label']}；可靠 {summary['reliable_count']} 条/"
        f"{fmt_duration_seconds(summary['reliable_seconds'])}，低置信 {summary['candidate_count']} 条/"
        f"{fmt_duration_seconds(summary['candidate_seconds'])}，丢弃 {summary['discarded_count']} 条/"
        f"{fmt_duration_seconds(summary['discarded_seconds'])}"
    )


def looks_repetitive(text: str) -> bool:
    text = strip_text(text)
    compact = re.sub(r"\s+", "", text)
    if len(compact) < 12:
        return False

    unique_ratio = len(set(compact)) / max(1, len(compact))
    if unique_ratio < 0.18:
        return True

    for size in range(1, min(12, len(compact) // 2) + 1):
        chunk = compact[:size]
        repeated = chunk * (len(compact) // size)
        if compact.startswith(repeated) and len(repeated) >= len(compact) * 0.8:
            return True

    phrases = re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,12}", compact)
    counts: Dict[str, int] = {}
    for phrase in phrases:
        counts[phrase] = counts.get(phrase, 0) + 1
        if counts[phrase] >= 5:
            return True
    return False


def ensure_ffmpeg(log: Callable[[str], None]) -> Path:
    try:
        import imageio_ffmpeg
    except ImportError as exc:
        raise RuntimeError("缺少 imageio-ffmpeg。请先运行 install_deps.bat 安装依赖。") from exc

    src = Path(imageio_ffmpeg.get_ffmpeg_exe())
    if not src.exists():
        raise RuntimeError("未找到 ffmpeg 可执行文件。")

    ffmpeg_dir = APP_DIR / ".runtime" / "ffmpeg-bin"
    ffmpeg_dir.mkdir(parents=True, exist_ok=True)
    dst = ffmpeg_dir / ("ffmpeg.exe" if os.name == "nt" else "ffmpeg")
    if not dst.exists():
        shutil.copy2(src, dst)
        log(f"已准备 ffmpeg：{dst}")
    os.environ["PATH"] = str(ffmpeg_dir) + os.pathsep + os.environ.get("PATH", "")
    return dst


def run_process(args: List[str], log: Callable[[str], None], timeout: Optional[int] = None) -> subprocess.CompletedProcess:
    log("运行：" + " ".join(f'"{x}"' if " " in x else x for x in args))
    return subprocess.run(
        args,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=timeout,
        check=False,
    )


def audio_duration(ffmpeg: Path, audio_path: Path) -> Optional[str]:
    proc = subprocess.run(
        [str(ffmpeg), "-hide_banner", "-i", str(audio_path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    match = re.search(r"Duration:\s*(\d+:\d+:\d+\.\d+)", proc.stderr)
    return match.group(1) if match else None


def enhance_audio(ffmpeg: Path, audio_path: Path, work_dir: Path, log: Callable[[str], None]) -> Path:
    work_dir.mkdir(parents=True, exist_ok=True)
    wav_path = work_dir / f"{clean_filename(audio_path.stem)}_enhanced.wav"
    filters = "highpass=f=120,lowpass=f=3800,afftdn=nf=-28,dynaudnorm=f=151:g=12:p=0.9"
    args = [
        str(ffmpeg),
        "-y",
        "-i",
        str(audio_path),
        "-ac",
        "1",
        "-ar",
        "16000",
        "-af",
        filters,
        str(wav_path),
    ]
    proc = run_process(args, log)
    if proc.returncode != 0:
        raise RuntimeError("音频增强失败：\n" + proc.stderr[-1500:])
    log(f"增强音频已生成：{wav_path}")
    return wav_path


def import_whisper():
    try:
        import whisper
    except ImportError as exc:
        raise RuntimeError("缺少 openai-whisper。请先运行 install_deps.bat 安装依赖。") from exc
    return whisper


def classify_segments(raw_segments: Iterable[dict], strict: bool) -> Tuple[List[dict], List[dict], List[dict], List[dict]]:
    reliable: List[dict] = []
    candidates: List[dict] = []
    discarded: List[dict] = []
    all_segments: List[dict] = []
    seen_text: Dict[str, int] = {}

    avg_limit = -1.0 if strict else -1.25
    no_speech_limit = 0.58 if strict else 0.72

    for seg in raw_segments:
        text = strip_text(seg.get("text", ""))
        if not text:
            continue

        key = re.sub(r"\s+", "", text)
        seen_text[key] = seen_text.get(key, 0) + 1

        avg_logprob = seg.get("avg_logprob")
        no_speech_prob = seg.get("no_speech_prob")
        compression_ratio = seg.get("compression_ratio")

        flags = []
        if avg_logprob is not None and avg_logprob < avg_limit:
            flags.append("平均置信偏低")
        if no_speech_prob is not None and no_speech_prob > no_speech_limit:
            flags.append("疑似非语音")
        if compression_ratio is not None and compression_ratio > 2.4:
            flags.append("疑似重复/压缩异常")
        repetitive = looks_repetitive(text)
        if repetitive:
            flags.append("文本重复异常")
        if seen_text[key] >= 3 and len(key) >= 4:
            flags.append("同句反复出现")

        record = {
            "start": float(seg.get("start", 0) or 0),
            "end": float(seg.get("end", 0) or 0),
            "text": text,
            "avg_logprob": avg_logprob,
            "no_speech_prob": no_speech_prob,
            "compression_ratio": compression_ratio,
            "flags": flags,
        }
        all_segments.append(record)
        discard_as_hallucination = repetitive or (
            len(key) >= 40
            and avg_logprob is not None
            and avg_logprob < avg_limit
            and no_speech_prob is not None
            and no_speech_prob > 0.5
        )
        if discard_as_hallucination:
            discarded.append(record)
        elif flags:
            candidates.append(record)
        else:
            reliable.append(record)

    return reliable, candidates, discarded, all_segments


def write_txt(path: Path, title: str, segments: List[dict], empty_text: str) -> None:
    lines = [title, "=" * len(title), ""]
    if segments:
        for seg in segments:
            lines.append(f"[{fmt_time(seg['start'])} - {fmt_time(seg['end'])}] {seg['text']}")
    else:
        lines.append(empty_text)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_srt(path: Path, segments: List[dict]) -> None:
    lines: List[str] = []
    for idx, seg in enumerate(segments, start=1):
        lines.append(str(idx))
        lines.append(f"{fmt_srt_time(seg['start'])} --> {fmt_srt_time(seg['end'])}")
        lines.append(seg["text"])
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def set_run_font(run, size: Optional[int] = None, bold: Optional[bool] = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def write_docx(
    path: Path,
    source_file: Path,
    duration: Optional[str],
    model_name: str,
    enhanced: bool,
    strict: bool,
    summary: Dict[str, Any],
    reliable: List[dict],
    candidates: List[dict],
) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx。请先运行 install_deps.bat 安装依赖。") from exc

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph()
    run = title.add_run("语音转文字结果")
    set_run_font(run, size=18, bold=True)

    meta = [
        f"源文件：{source_file.name}",
        f"音频时长：{duration or '未读取到'}",
        f"模型：Whisper {model_name}",
        f"处理：{'降噪增强；' if enhanced else ''}{'严格过滤低置信片段' if strict else '普通过滤'}",
        f"质量摘要：{format_quality_summary(summary)}",
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    for item in meta:
        doc.add_paragraph(item)

    heading = doc.add_paragraph()
    set_run_font(heading.add_run("正式转写正文"), size=14, bold=True)
    if reliable:
        for seg in reliable:
            doc.add_paragraph(f"{fmt_time(seg['start'])} - {fmt_time(seg['end'])}：{seg['text']}")
    else:
        p = doc.add_paragraph()
        r = p.add_run("未检测到可靠、连续、可确认的中文语音内容。")
        set_run_font(r, bold=True)

    heading = doc.add_paragraph()
    set_run_font(heading.add_run("低置信候选片段（仅供人工核听参考）"), size=14, bold=True)
    if candidates:
        for seg in candidates:
            flags = "、".join(seg["flags"]) if seg["flags"] else "低置信"
            avg = seg.get("avg_logprob")
            ns = seg.get("no_speech_prob")
            avg_text = "" if avg is None else f"，平均置信 {avg:.2f}"
            ns_text = "" if ns is None else f"，非语音概率 {ns:.2f}"
            doc.add_paragraph(
                f"{fmt_time(seg['start'])} - {fmt_time(seg['end'])}：{seg['text']}（{flags}{avg_text}{ns_text}）"
            )
    else:
        doc.add_paragraph("无。")

    heading = doc.add_paragraph()
    set_run_font(heading.add_run("说明"), size=14, bold=True)
    doc.add_paragraph(
        "自动语音识别会受录音音量、噪声、距离、多人重叠说话和方言影响。若正式正文为空或候选片段较多，请以原录音人工核听为准。"
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def transcribe_file(
    audio_path: Path,
    output_dir: Path,
    model_name: str,
    enhance: bool,
    strict: bool,
    log: Callable[[str], None],
    loaded_model: Optional[Any] = None,
) -> Dict[str, Path]:
    if not audio_path.exists():
        raise FileNotFoundError(audio_path)

    output_dir.mkdir(parents=True, exist_ok=True)
    run_dir = output_dir / f"{clean_filename(audio_path.stem)}_{now_stamp()}"
    run_dir.mkdir(parents=True, exist_ok=True)

    ffmpeg = ensure_ffmpeg(log)
    duration = audio_duration(ffmpeg, audio_path)
    if duration:
        log(f"音频时长：{duration}")

    source_for_asr = audio_path
    if enhance:
        log("开始降噪增强音频。")
        source_for_asr = enhance_audio(ffmpeg, audio_path, run_dir, log)

    if loaded_model is None:
        whisper = import_whisper()
        log(f"加载 Whisper 模型：{model_name}。首次使用会下载模型，请耐心等待。")
        model = whisper.load_model(model_name)
        log("模型加载完成，开始识别。")
    else:
        model = loaded_model
        log("复用已加载模型，开始识别。")

    result = model.transcribe(
        str(source_for_asr),
        language="zh",
        task="transcribe",
        fp16=False,
        verbose=False,
        temperature=0.0,
        initial_prompt=None,
        condition_on_previous_text=False,
        no_speech_threshold=0.5 if strict else 0.6,
        logprob_threshold=-0.8 if strict else -1.0,
        compression_ratio_threshold=2.4,
    )

    reliable, candidates, discarded, all_segments = classify_segments(result.get("segments", []), strict=strict)
    log(f"识别完成：可靠片段 {len(reliable)} 条，低置信候选 {len(candidates)} 条，丢弃幻觉片段 {len(discarded)} 条。")
    summary = build_quality_summary(reliable, candidates, discarded, duration)
    log("质量摘要：" + format_quality_summary(summary))

    stem = clean_filename(audio_path.stem)
    raw_json = run_dir / f"{stem}_raw.json"
    report_json = run_dir / f"{stem}_report.json"
    transcript_txt = run_dir / f"{stem}_正式转写.txt"
    candidates_txt = run_dir / f"{stem}_低置信候选.txt"
    srt_path = run_dir / f"{stem}_正式转写.srt"
    docx_path = run_dir / f"{stem}_语音转文字结果.docx"

    raw_json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    report_json.write_text(
        json.dumps(
            {
                "source": str(audio_path),
                "duration": duration,
                "model": model_name,
                "enhance": enhance,
                "strict": strict,
                "summary": summary,
                "reliable": reliable,
                "candidates": candidates,
                "discarded": discarded,
                "all_segments": all_segments,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    write_txt(transcript_txt, "正式转写", reliable, "未检测到可靠、连续、可确认的中文语音内容。")
    write_txt(candidates_txt, "低置信候选片段（仅供人工核听参考）", candidates, "无低置信候选片段。")
    write_srt(srt_path, reliable)
    write_docx(docx_path, audio_path, duration, model_name, enhance, strict, summary, reliable, candidates)

    outputs = {
        "run_dir": run_dir,
        "docx": docx_path,
        "txt": transcript_txt,
        "candidates": candidates_txt,
        "srt": srt_path,
        "report": report_json,
        "raw": raw_json,
    }
    log(f"输出目录：{run_dir}")
    log(f"Word：{docx_path}")
    return outputs


def transcribe_batch(
    audio_paths: Sequence[Path],
    output_dir: Path,
    model_name: str,
    enhance: bool,
    strict: bool,
    log: Callable[[str], None],
) -> Dict[str, Any]:
    if not audio_paths:
        raise ValueError("没有可转写的音频文件。")

    output_dir.mkdir(parents=True, exist_ok=True)
    batch_stamp = now_stamp()
    log(f"批量任务开始：{len(audio_paths)} 个文件。")
    whisper = import_whisper()
    log(f"加载 Whisper 模型：{model_name}。首次使用会下载模型，请耐心等待。")
    model = whisper.load_model(model_name)
    log("模型加载完成。")

    items: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    for index, audio_path in enumerate(audio_paths, start=1):
        log(f"处理 {index}/{len(audio_paths)}：{audio_path}")
        if not is_supported_audio(audio_path):
            log(f"提示：{audio_path.name} 的扩展名不在常见支持列表内，将继续尝试。")
        try:
            outputs = transcribe_file(
                audio_path=audio_path,
                output_dir=output_dir,
                model_name=model_name,
                enhance=enhance,
                strict=strict,
                log=log,
                loaded_model=model,
            )
            items.append(
                {
                    "source": str(audio_path),
                    "status": "ok",
                    "run_dir": str(outputs["run_dir"]),
                    "docx": str(outputs["docx"]),
                    "txt": str(outputs["txt"]),
                    "candidates": str(outputs["candidates"]),
                    "srt": str(outputs["srt"]),
                    "report": str(outputs["report"]),
                    "raw": str(outputs["raw"]),
                }
            )
        except Exception as exc:
            message = str(exc)
            log(f"处理失败：{audio_path}；{message}")
            error_item = {"source": str(audio_path), "status": "error", "error": message}
            items.append(error_item)
            errors.append(error_item)

    manifest = output_dir / f"batch_manifest_{batch_stamp}.json"
    manifest.write_text(
        json.dumps(
            {
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "model": model_name,
                "enhance": enhance,
                "strict": strict,
                "total": len(audio_paths),
                "succeeded": len(audio_paths) - len(errors),
                "failed": len(errors),
                "items": items,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    log(f"批量清单：{manifest}")
    return {"manifest": manifest, "items": items, "errors": errors}
